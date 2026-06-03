import os
from docx import Document

garantias_dir = r'c:\centromotosWeb\backend\garantias'

# Analizar cada archivo Word
for filename in sorted(os.listdir(garantias_dir)):
    if filename.endswith('.docx'):
        filepath = os.path.join(garantias_dir, filename)
        doc = Document(filepath)
        
        print(f'\n{"="*60}')
        print(f'ARCHIVO: {filename}')
        print(f'{"="*60}')
        
        # Mostrar todos los párrafos
        for i, parrafo in enumerate(doc.paragraphs):
            texto = parrafo.text.strip()
            if texto:
                print(f'P{i}: {repr(texto)}')
        
        # Mostrar tablas
        if doc.tables:
            print(f'\nTABLAS: {len(doc.tables)}')
            for table_idx, table in enumerate(doc.tables):
                for row_idx, row in enumerate(table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        texto = cell.text.strip()
                        if texto:
                            print(f'T{table_idx}[{row_idx},{col_idx}]: {repr(texto)}')
